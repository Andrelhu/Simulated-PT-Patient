import io
import json
import secrets
from datetime import datetime, timezone
from functools import wraps
from flask import (Flask, request, jsonify, render_template_string, session,
                   redirect, send_file, Response)
from pathlib import Path
from werkzeug.security import generate_password_hash, check_password_hash

# --- Configuration ---
API_URL    = "https://carc.harrisburgu.edu/api/v1/projects/vm-for-r-projects/llm"
API_KEY    = open(Path.home() / "apikey.txt").read().strip()
MODEL_NAME = "Gemma 4 E4B"

CHARS_DIR    = Path(__file__).parent
SESSIONS_DIR = CHARS_DIR / "sessions"
AVATARS_DIR  = CHARS_DIR / "avatars"
SESSIONS_DIR.mkdir(exist_ok=True)
AVATARS_DIR.mkdir(exist_ok=True)
USERS_FILE   = CHARS_DIR / "users.json"
META_FILE    = CHARS_DIR / "character_meta.json"

ROLES = ("student", "teacher", "developer")

# edge-tts voices used when a character has no explicit "voice" in
# character_meta.json. Run `edge-tts --list-voices` to see the full catalogue.
DEFAULT_VOICES = {"female": "en-US-AriaNeural", "male": "en-US-GuyNeural"}
MAX_TTS_CHARS  = 4000


# ── User helpers ──────────────────────────────────────────────────────────────
def load_users():
    """Returns {username: {password, role, created_at}}.
    Migrates the old flat {username: password_hash} format transparently."""
    if not USERS_FILE.exists():
        return {}
    try:
        raw = json.loads(USERS_FILE.read_text(encoding="utf-8"))
    except Exception:
        return {}
    users = {}
    for name, val in raw.items():
        if isinstance(val, str):
            users[name] = {"password": val, "role": "student", "created_at": ""}
        else:
            val.setdefault("role", "student")
            val.setdefault("created_at", "")
            users[name] = val
    return users


def save_users(users):
    USERS_FILE.write_text(json.dumps(users, indent=2, ensure_ascii=False), encoding="utf-8")


def current_role():
    return load_users().get(session.get("username", ""), {}).get("role", "student")


# ── Character helpers ─────────────────────────────────────────────────────────
def character_names():
    return sorted(f.stem.replace("character_", "", 1)
                  for f in CHARS_DIR.glob("character_*.txt"))


def load_meta():
    if META_FILE.exists():
        try:
            return json.loads(META_FILE.read_text(encoding="utf-8"))
        except Exception:
            pass
    return {}


def read_character(name):
    """Reads a character's system context. Returns '' for unknown names."""
    if name not in character_names():
        return ""
    return (CHARS_DIR / f"character_{name}.txt").read_text(encoding="utf-8")


def build_avatar_svg(name):
    """Deterministic colored silhouette, used when no image file exists."""
    hue = sum(ord(c) for c in name) % 360
    return f"""<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 200 200" width="200" height="200">
  <defs>
    <linearGradient id="bg" x1="0" y1="0" x2="0" y2="1">
      <stop offset="0" stop-color="hsl({hue}, 42%, 74%)"/>
      <stop offset="1" stop-color="hsl({hue}, 38%, 52%)"/>
    </linearGradient>
  </defs>
  <rect width="200" height="200" fill="url(#bg)"/>
  <circle cx="100" cy="76" r="33" fill="rgba(255,255,255,0.88)"/>
  <ellipse cx="100" cy="178" rx="55" ry="46" fill="rgba(255,255,255,0.88)"/>
</svg>"""


# ── Session helpers ───────────────────────────────────────────────────────────
def _session_path(username, session_id):
    return SESSIONS_DIR / username / f"session_{session_id}.json"


def save_session(session_id, username, character, exchanges):
    user_dir = SESSIONS_DIR / username
    user_dir.mkdir(exist_ok=True)
    path = _session_path(username, session_id)
    now = datetime.now(timezone.utc).isoformat()
    started_at, notes = now, ""
    if path.exists():
        try:
            prev = json.loads(path.read_text(encoding="utf-8"))
            started_at = prev.get("started_at", now)
            notes      = prev.get("notes", "")
        except Exception:
            pass
    data = {
        "session_id":     session_id,
        "username":       username,
        "character":      character,
        "started_at":     started_at,
        "last_updated":   now,
        "exchange_count": len(exchanges),
        "notes":          notes,
        "exchanges":      [{"question": q, "response": r} for q, r in exchanges],
    }
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def get_user_sessions(username):
    user_dir = SESSIONS_DIR / username
    if not user_dir.exists():
        return []
    result = []
    for p in sorted(user_dir.glob("session_*.json"), key=lambda x: x.stat().st_mtime, reverse=True):
        try:
            d = json.loads(p.read_text(encoding="utf-8"))
            result.append({
                "session_id":     d.get("session_id", p.stem.replace("session_", "")),
                "character":      d.get("character", "Unknown"),
                "started_at":     d.get("started_at", "")[:19].replace("T", " "),
                "exchange_count": d.get("exchange_count", 0),
                "has_notes":      bool(d.get("notes", "").strip()),
            })
        except Exception:
            continue
    return result


def get_session_data(username, session_id):
    path = _session_path(username, session_id)
    if not path.exists():
        return None
    try:
        d = json.loads(path.read_text(encoding="utf-8"))
        d["started_at"] = d.get("started_at", "")[:19].replace("T", " ")
        d.setdefault("notes", "")
        return d
    except Exception:
        return None


# ── Main chat UI ──────────────────────────────────────────────────────────────
HTML = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Simulated PT Patient</title>
<style>
  * { box-sizing: border-box; margin: 0; padding: 0; }
  body { font-family: system-ui, sans-serif; background: #f0f2f5;
         display: flex; flex-direction: column; height: 100vh; overflow: hidden; }

  /* ── header ── */
  header { background: #0C6157; color: white; padding: 10px 20px;
           display: flex; align-items: center; justify-content: space-between; flex-shrink: 0; }
  #header-logo { height: 44px; width: auto; flex-shrink: 0; }
  #header-center { flex: 1; text-align: center; }
  header h1 { font-size: 1.05rem; font-weight: 600; }
  header p  { font-size: 0.78rem; opacity: 0.85; margin-top: 2px; }
  #header-actions { display: flex; gap: 6px; align-items: center; flex-shrink: 0; }
  #username-display { font-size: 0.75rem; opacity: 0.85; margin-right: 4px; white-space: nowrap; }
  .role-badge { background: rgba(255,255,255,.2); border-radius: 4px;
                padding: 1px 6px; font-size: 0.68rem; margin-left: 4px;
                text-transform: uppercase; letter-spacing: .04em; }
  .hdr-btn { background: #CBB778; border: none; color: white; border-radius: 8px;
             padding: 6px 12px; cursor: pointer; font-size: 0.82rem; white-space: nowrap;
             text-decoration: none; display: inline-block; }
  .hdr-btn:hover { background: #b5a265; }
  #test-btn.running { background: rgba(220,50,50,0.85); }
  #test-btn.running:hover { background: rgba(220,50,50,1); }
  #tts-btn.active { background: #0C6157; box-shadow: inset 0 0 0 2px #CBB778; }
  #test-progress { background: #e8f0fe; color: #1a73e8; font-size: 0.8rem;
                   padding: 6px 16px; text-align: center; flex-shrink: 0;
                   display: none; border-bottom: 1px solid #c5d4fb; }

  /* ── layout ── */
  #app-layout { display: flex; flex: 1; overflow: hidden; }

  /* ── sidebar (1/3 of window) ── */
  /* Exactly one third of the viewport. No max-width — that was capping it
     to about a quarter on wide screens. */
  #sidebar { flex: 0 0 33.333%; width: 33.333%; min-width: 260px;
             background: white; border-right: 1px solid #e0e0e0;
             display: flex; flex-direction: column; padding: 20px 18px; overflow: hidden; }
  #avatar-img { width: 100%; max-width: 340px; aspect-ratio: 1 / 1; object-fit: cover;
                border-radius: 16px; display: block; margin: 0 auto 14px;
                background: #eef3f2; }
  #sidebar-char-name { text-align: center; font-size: 1.25rem; font-weight: 600;
                       color: #0C6157; margin-bottom: 12px; word-break: break-word; }
  #char-select { width: 100%; padding: 8px 10px; border: 1px solid #ccc;
                 border-radius: 8px; font-size: 0.9rem; margin-bottom: 16px; }
  .sidebar-divider { border: none; border-top: 1px solid #eee; margin-bottom: 14px; }
  #notes-wrap { display: flex; flex-direction: column; flex: 1; min-height: 0; }
  .notes-head { display: flex; justify-content: space-between; align-items: baseline;
                margin-bottom: 8px; }
  .notes-label { font-size: 0.68rem; font-weight: 700; color: #aaa;
                 text-transform: uppercase; letter-spacing: 0.06em; }
  #notes-status { font-size: 0.7rem; color: #388e3c; }
  #notes-area { flex: 1; min-height: 120px; width: 100%; padding: 10px 12px;
                border: 1px solid #ddd; border-radius: 8px; font-size: 0.85rem;
                font-family: inherit; line-height: 1.5; resize: none; outline: none; }
  #notes-area:focus { border-color: #CBB778; }

  /* ── main column ── */
  #main-col { display: flex; flex-direction: column; flex: 1; overflow: hidden; }

  /* ── settings panel (developer only) ── */
  #settings { background: #fff; border-bottom: 1px solid #ddd; padding: 14px 20px;
              display: none; flex-shrink: 0; gap: 10px; flex-direction: column; }
  #settings.open { display: flex; }
  #settings label { font-size: 0.82rem; color: #555; font-weight: 500;
                    display: flex; flex-direction: column; gap: 4px; }
  #ctx-area { width: 100%; height: 190px; padding: 8px 10px; border: 1px solid #ccc;
              border-radius: 8px; font-size: 0.82rem; font-family: monospace; resize: vertical; }
  .settings-actions { display: flex; gap: 8px; align-items: center; }
  .btn-sm { padding: 7px 16px; font-size: 0.85rem; border-radius: 8px;
            border: none; cursor: pointer; background: #CBB778; color: white; }
  .btn-sm:hover { background: #b5a265; }
  #status-msg { font-size: 0.8rem; color: #388e3c; }

  /* ── chat ── */
  #chat { flex: 1; overflow-y: auto; padding: 16px; display: flex; flex-direction: column; gap: 10px; }
  .bubble { max-width: 78%; padding: 10px 14px; border-radius: 18px;
            line-height: 1.5; font-size: 0.95rem; white-space: pre-wrap; }
  .user  { background: #1a73e8; color: white; align-self: flex-end; border-bottom-right-radius: 4px; }
  .agent { background: white; color: #111; align-self: flex-start;
           border-bottom-left-radius: 4px; box-shadow: 0 1px 2px rgba(0,0,0,.1); }
  .label { font-size: 0.72rem; color: #888; margin-bottom: 2px; }
  .user-wrap  { align-self: flex-end;  display: flex; flex-direction: column; align-items: flex-end; }
  .agent-wrap { align-self: flex-start; display: flex; flex-direction: column; }
  .typing { color: #888; font-style: italic; font-size: 0.9rem; padding: 4px 14px; }

  /* ── input bar ── */
  #input-bar { display: flex; gap: 8px; padding: 12px 16px;
               background: white; border-top: 1px solid #ddd; flex-shrink: 0; }
  #msg { flex: 1; padding: 10px 14px; border: 1px solid #ccc; border-radius: 24px;
         font-size: 0.95rem; outline: none; }
  #msg:focus { border-color: #CBB778; }
  #send { background: #CBB778; color: white; border: none; border-radius: 24px;
          padding: 10px 20px; cursor: pointer; font-size: 0.95rem; }
  #send:disabled { opacity: 0.5; cursor: default; }

  @media (max-width: 820px) {
    #sidebar { flex: 0 0 240px; min-width: 200px; padding: 14px; }
    #avatar-img { max-width: 150px; }
    #sidebar-char-name { font-size: 1rem; }
    header p { display: none; }
  }
</style>
</head>
<body>

<header>
  <img id="header-logo" src="https://www.arcgis.com/sharing/rest/content/items/088d68905927400bb34449dc1b387446/resources/images/widget_2/1709839675447.png" alt="logo">
  <div id="header-center">
    <h1 id="header-title">Simulated PT Patient</h1>
    <p>Conduct yourself as you would in a real clinical setting.</p>
  </div>
  <div id="header-actions">
    <span id="username-display">{{ username }}{% if role != 'student' %}<span class="role-badge">{{ role }}</span>{% endif %}</span>
    <button id="new-chat-btn" class="hdr-btn">&#43; New</button>
    <button id="tts-btn" class="hdr-btn">&#128264; Voice</button>
    {% if role == 'developer' %}
    <button id="test-btn" class="hdr-btn">&#9654; Test</button>
    <button id="toggle-settings" class="hdr-btn">&#9881; Settings</button>
    {% endif %}
    <a href="/sessions" class="hdr-btn">Sessions</a>
    <a href="/logout" class="hdr-btn">Logout</a>
  </div>
</header>
<div id="test-progress"></div>

<div id="app-layout">

  <aside id="sidebar">
    <img id="avatar-img" src="" alt="patient avatar">
    <div id="sidebar-char-name">&mdash;</div>
    <select id="char-select"></select>
    <hr class="sidebar-divider">
    <div id="notes-wrap">
      <div class="notes-head">
        <span class="notes-label">Scrap Notes</span>
        <span id="notes-status"></span>
      </div>
      <textarea id="notes-area" placeholder="Jot down findings, hypotheses, follow-up questions…"></textarea>
    </div>
  </aside>

  <div id="main-col">
    {% if role == 'developer' %}
    <div id="settings">
      <label>
        System context (developer view &mdash; edits apply to this chat only)
        <textarea id="ctx-area"></textarea>
      </label>
      <div class="settings-actions">
        <button class="btn-sm" id="apply-btn">Apply &amp; reset chat</button>
        <button class="btn-sm" id="reload-ctx-btn">Reload from file</button>
        <span id="status-msg"></span>
      </div>
    </div>
    {% endif %}

    <div id="chat"></div>

    <div id="input-bar">
      <input id="msg" type="text" placeholder="Type your question…" autocomplete="off">
      <button id="send">Send</button>
    </div>
  </div>
</div>

<script>
  const IS_DEV = {{ 'true' if role == 'developer' else 'false' }};

  const chat       = document.getElementById('chat');
  const msgInput   = document.getElementById('msg');
  const sendBtn    = document.getElementById('send');
  const charSelect = document.getElementById('char-select');
  const notesArea  = document.getElementById('notes-area');
  const notesStatus= document.getElementById('notes-status');
  const ctxArea    = document.getElementById('ctx-area');      // null for non-devs
  const statusMsg  = document.getElementById('status-msg');    // null for non-devs

  let history          = [];
  let sessionId        = crypto.randomUUID();
  let currentCharacter = '';
  let currentGender    = 'female';

  function setStatus(txt, ms) {
    if (!statusMsg) return;
    statusMsg.textContent = txt;
    if (ms) setTimeout(() => { statusMsg.textContent = ''; }, ms);
  }

  // ── Text to speech ──
  const synth = window.speechSynthesis;
  let voices  = [];
  let ttsOn   = false;
  const ttsBtn = document.getElementById('tts-btn');

  function refreshVoices() { voices = synth ? synth.getVoices() : []; }
  refreshVoices();
  if (synth && synth.onvoiceschanged !== undefined) synth.onvoiceschanged = refreshVoices;

  const FEMALE_HINTS = ['zira','aria','jenny','michelle','samantha','victoria','karen',
                        'moira','tessa','fiona','serena','allison','joanna','female','woman'];
  const MALE_HINTS   = ['david','mark','guy','christopher','eric','roger','steffan','alex',
                        'daniel','fred','oliver','thomas','brian','male','man'];

  function pickVoice(gender) {
    const english = voices.filter(v => /^en/i.test(v.lang));
    const pool    = english.length ? english : voices;
    const hints   = gender === 'male' ? MALE_HINTS : FEMALE_HINTS;
    for (const hint of hints) {
      const match = pool.find(voice => voice.name.toLowerCase().includes(hint));
      if (match) return match;
    }
    return pool[0] || null;
  }

  ttsBtn.addEventListener('click', () => {
    ttsOn = !ttsOn;
    ttsBtn.classList.toggle('active', ttsOn);
    ttsBtn.textContent = ttsOn ? '🔊 Voice' : '🔈 Voice';
    if (!ttsOn) stopSpeaking();
  });

  // Server-side neural voices (edge-tts) with the browser's own voices as a
  // fallback, so the app still speaks if the VM loses network.
  let currentAudio = null;
  let speakToken   = 0;

  function stopSpeaking() {
    speakToken++;
    if (synth) synth.cancel();
    if (currentAudio) {
      currentAudio.pause();
      URL.revokeObjectURL(currentAudio.src);
      currentAudio = null;
    }
  }

  function speakBrowser(text) {
    if (!synth) return;
    const utt = new SpeechSynthesisUtterance(text);
    const voice = pickVoice(currentGender);
    if (voice) utt.voice = voice;
    utt.rate  = 0.95;
    utt.pitch = currentGender === 'male' ? 0.9 : 1.05;
    synth.speak(utt);
  }

  async function speak(raw) {
    if (!ttsOn || !raw) return;
    const text = raw.replace(/[*_#`]/g, '').trim();
    if (!text) return;

    stopSpeaking();
    const token = speakToken;

    try {
      const res = await fetch('/tts', {
        method: 'POST',
        headers: { 'Content-Type': 'application/json' },
        body: JSON.stringify({ character: currentCharacter, text: text })
      });
      if (!res.ok) throw new Error('tts unavailable');
      const blob = await res.blob();
      if (token !== speakToken) return;   // superseded while we waited

      const audio = new Audio(URL.createObjectURL(blob));
      currentAudio = audio;
      audio.addEventListener('ended', () => { URL.revokeObjectURL(audio.src); });
      await audio.play();
    } catch (e) {
      if (token === speakToken) speakBrowser(text);
    }
  }

  // ── Scrap notes ──
  let notesTimer = null;
  notesArea.addEventListener('input', () => {
    clearTimeout(notesTimer);
    notesStatus.textContent = '';
    notesTimer = setTimeout(saveNotes, 800);
  });

  async function saveNotes() {
    try {
      await fetch('/session-notes', {
        method: 'POST',
        headers: { 'Content-Type': 'application/json' },
        body: JSON.stringify({ session_id: sessionId, character: currentCharacter,
                               notes: notesArea.value })
      });
      notesStatus.textContent = 'Saved';
      setTimeout(() => { notesStatus.textContent = ''; }, 1500);
    } catch (e) {
      notesStatus.textContent = 'Not saved';
    }
  }

  // ── Characters ──
  async function loadCharacterList() {
    const res  = await fetch('/characters');
    const data = await res.json();
    charSelect.innerHTML = '';
    data.characters.forEach(c => {
      const opt = document.createElement('option');
      opt.value = c.name;
      opt.textContent = c.name;
      opt.dataset.gender = c.gender || 'female';
      charSelect.appendChild(opt);
    });
    if (data.characters.length) await selectCharacter(data.characters[0].name);
  }

  async function selectCharacter(name) {
    currentCharacter = name;
    charSelect.value = name;
    const opt = Array.from(charSelect.options).find(o => o.value === name);
    currentGender = opt ? (opt.dataset.gender || 'female') : 'female';

    document.getElementById('sidebar-char-name').textContent = name;
    document.getElementById('header-title').textContent = name + ' — Simulated PT Patient';
    document.getElementById('avatar-img').src = '/avatar/' + encodeURIComponent(name);

    if (IS_DEV && ctxArea) {
      const r = await fetch('/character/' + encodeURIComponent(name));
      const d = await r.json();
      ctxArea.value = d.context || '';
    }
    startNewChat();
  }

  charSelect.addEventListener('change', () => selectCharacter(charSelect.value));

  // ── New chat ──
  function startNewChat() {
    history   = [];
    sessionId = crypto.randomUUID();
    chat.innerHTML = '';
    notesArea.value = '';
    notesStatus.textContent = '';
    stopSpeaking();
  }

  document.getElementById('new-chat-btn').addEventListener('click', () => {
    startNewChat();
    msgInput.focus();
  });

  if (IS_DEV && ctxArea) {
    document.getElementById('apply-btn').addEventListener('click', () => {
      startNewChat();
      setStatus('Context applied. Chat reset.', 2500);
      document.getElementById('settings').classList.remove('open');
      msgInput.focus();
    });
    document.getElementById('reload-ctx-btn').addEventListener('click', async () => {
      const r = await fetch('/character/' + encodeURIComponent(currentCharacter));
      const d = await r.json();
      ctxArea.value = d.context || '';
      setStatus('Reloaded from file.', 2000);
    });
    document.getElementById('toggle-settings').addEventListener('click', () => {
      document.getElementById('settings').classList.toggle('open');
    });
  }

  // ── Chat ──
  function addBubble(role, text) {
    const wrap = document.createElement('div');
    wrap.className = role === 'user' ? 'user-wrap' : 'agent-wrap';
    const label = document.createElement('div');
    label.className = 'label';
    label.textContent = role === 'user' ? 'You' : (currentCharacter || 'Patient');
    const bubble = document.createElement('div');
    bubble.className = 'bubble ' + (role === 'user' ? 'user' : 'agent');
    bubble.textContent = text;
    wrap.appendChild(label);
    wrap.appendChild(bubble);
    chat.appendChild(wrap);
    chat.scrollTop = chat.scrollHeight;
  }

  function payloadFor(text) {
    return JSON.stringify({
      message: text,
      history: history,
      session_id: sessionId,
      character: currentCharacter,
      override_context: (IS_DEV && ctxArea) ? ctxArea.value.trim() : ''
    });
  }

  async function sendMsg() {
    const text = msgInput.value.trim();
    if (!text || !currentCharacter) return;
    msgInput.value = '';
    sendBtn.disabled = true;
    addBubble('user', text);

    const typing = document.createElement('div');
    typing.className = 'typing';
    typing.textContent = (currentCharacter || 'Patient') + ' is typing…';
    chat.appendChild(typing);
    chat.scrollTop = chat.scrollHeight;

    try {
      const res  = await fetch('/chat', {
        method: 'POST',
        headers: { 'Content-Type': 'application/json' },
        body: payloadFor(text)
      });
      const data = await res.json();
      typing.remove();
      addBubble('agent', data.reply);
      history.push([text, data.reply]);
      speak(data.reply);
    } catch (e) {
      typing.remove();
      addBubble('agent', 'Connection error — please try again.');
    }
    sendBtn.disabled = false;
    msgInput.focus();
  }

  sendBtn.addEventListener('click', sendMsg);
  msgInput.addEventListener('keydown', e => { if (e.key === 'Enter') sendMsg(); });

  // ── Test runner (developer only) ──
  const testBtn      = document.getElementById('test-btn');
  const testProgress = document.getElementById('test-progress');
  let   stopRequested = false;

  async function runTest() {
    if (testBtn.classList.contains('running')) { stopRequested = true; return; }
    if (!currentCharacter) return;

    startNewChat();
    stopRequested = false;
    testBtn.textContent = '■ Stop';
    testBtn.classList.add('running');
    sendBtn.disabled  = true;
    msgInput.disabled = true;
    testProgress.style.display = 'block';

    const res       = await fetch('/test-questions');
    const data      = await res.json();
    const questions = data.questions;

    for (let i = 0; i < questions.length; i++) {
      if (stopRequested) { testProgress.textContent = `Stopped at question ${i + 1}.`; break; }
      const q = questions[i];
      testProgress.textContent = `Test running — question ${i + 1} of ${questions.length}`;
      addBubble('user', q);

      const typing = document.createElement('div');
      typing.className = 'typing';
      typing.textContent = (currentCharacter || 'Patient') + ' is typing…';
      chat.appendChild(typing);
      chat.scrollTop = chat.scrollHeight;

      try {
        const r    = await fetch('/chat', {
          method: 'POST',
          headers: { 'Content-Type': 'application/json' },
          body: payloadFor(q)
        });
        const resp = await r.json();
        typing.remove();
        addBubble('agent', resp.reply);
        history.push([q, resp.reply]);
        speak(resp.reply);
      } catch (e) {
        typing.remove();
        addBubble('agent', 'Connection error — test aborted.');
        break;
      }
    }

    if (!stopRequested) testProgress.textContent = `Test complete — ${questions.length} questions answered.`;
    setTimeout(() => { testProgress.style.display = 'none'; }, 4000);
    testBtn.innerHTML = '&#9654; Test';
    testBtn.classList.remove('running');
    sendBtn.disabled  = false;
    msgInput.disabled = false;
    msgInput.focus();
  }

  if (testBtn) testBtn.addEventListener('click', runTest);

  // ── Init ──
  loadCharacterList();
</script>
</body>
</html>"""


# ── Login page ────────────────────────────────────────────────────────────────
LOGIN_HTML = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Login — Simulated PT Patient</title>
<style>
  * { box-sizing: border-box; margin: 0; padding: 0; }
  body { font-family: system-ui, sans-serif; background: #f0f2f5;
         display: flex; align-items: center; justify-content: center; height: 100vh; }
  .card { background: white; border-radius: 12px; padding: 40px 36px;
          box-shadow: 0 2px 16px rgba(0,0,0,.12); width: 100%; max-width: 360px; }
  .logo { display: block; height: 56px; margin: 0 auto 20px; }
  h1 { text-align: center; font-size: 1.1rem; color: #0C6157; margin-bottom: 6px; }
  p  { text-align: center; font-size: 0.82rem; color: #777; margin-bottom: 24px; }
  input { width: 100%; padding: 10px 14px; border: 1px solid #ccc; border-radius: 8px;
          font-size: 0.95rem; margin-bottom: 14px; outline: none; }
  input:focus { border-color: #CBB778; }
  button { width: 100%; padding: 11px; background: #CBB778; color: white; border: none;
           border-radius: 8px; font-size: 1rem; cursor: pointer; margin-bottom: 12px; }
  button:hover { background: #b5a265; }
  .error { color: #c0392b; font-size: 0.83rem; margin-bottom: 10px; text-align: center; }
  .alt-link { text-align: center; font-size: 0.83rem; color: #555; }
  .alt-link a { color: #0C6157; text-decoration: none; font-weight: 500; }
  .alt-link a:hover { text-decoration: underline; }
</style>
</head>
<body>
<div class="card">
  <img class="logo" src="https://www.arcgis.com/sharing/rest/content/items/088d68905927400bb34449dc1b387446/resources/images/widget_2/1709839675447.png" alt="logo">
  <h1>Simulated PT Patient</h1>
  <p>Sign in to continue.</p>
  {% if error %}<div class="error">{{ error }}</div>{% endif %}
  <form method="POST">
    <input type="text" name="username" placeholder="Username" autofocus autocomplete="username">
    <input type="password" name="password" placeholder="Password" autocomplete="current-password">
    <button type="submit">Sign In</button>
  </form>
  <div class="alt-link">New user? <a href="/register">Create an account</a></div>
</div>
</body>
</html>"""


# ── Registration page ─────────────────────────────────────────────────────────
REGISTER_HTML = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Register — Simulated PT Patient</title>
<style>
  * { box-sizing: border-box; margin: 0; padding: 0; }
  body { font-family: system-ui, sans-serif; background: #f0f2f5;
         display: flex; align-items: center; justify-content: center;
         min-height: 100vh; padding: 24px; }
  .card { background: white; border-radius: 12px; padding: 36px;
          box-shadow: 0 2px 16px rgba(0,0,0,.12); width: 100%; max-width: 480px; }
  .logo { display: block; height: 48px; margin: 0 auto 16px; }
  h1 { text-align: center; font-size: 1.1rem; color: #0C6157; margin-bottom: 4px; }
  .subtitle { text-align: center; font-size: 0.82rem; color: #777; margin-bottom: 22px; }
  .consent-box { background: #f8f9fa; border: 1px solid #ddd; border-radius: 8px;
                 padding: 14px 16px; font-size: 0.8rem; color: #444; line-height: 1.6;
                 margin-bottom: 16px; max-height: 200px; overflow-y: auto; }
  .consent-box h3 { font-size: 0.85rem; color: #0C6157; margin-bottom: 8px; }
  .consent-box ol { padding-left: 18px; }
  .consent-box li { margin-bottom: 6px; }
  input[type=text], input[type=password] {
    width: 100%; padding: 10px 14px; border: 1px solid #ccc; border-radius: 8px;
    font-size: 0.95rem; margin-bottom: 12px; outline: none; }
  input:focus { border-color: #CBB778; }
  .checkbox-row { display: flex; align-items: flex-start; gap: 10px; margin-bottom: 18px;
                  font-size: 0.82rem; color: #444; cursor: pointer; }
  .checkbox-row input { width: auto; margin: 0; margin-top: 2px; cursor: pointer; }
  button { width: 100%; padding: 11px; background: #CBB778; color: white; border: none;
           border-radius: 8px; font-size: 1rem; cursor: pointer; margin-bottom: 12px; }
  button:hover { background: #b5a265; }
  .error { color: #c0392b; font-size: 0.83rem; margin-bottom: 12px; text-align: center; }
  .alt-link { text-align: center; font-size: 0.83rem; color: #555; }
  .alt-link a { color: #0C6157; text-decoration: none; font-weight: 500; }
  .alt-link a:hover { text-decoration: underline; }
</style>
</head>
<body>
<div class="card">
  <img class="logo" src="https://www.arcgis.com/sharing/rest/content/items/088d68905927400bb34449dc1b387446/resources/images/widget_2/1709839675447.png" alt="logo">
  <h1>Create Account</h1>
  <p class="subtitle">Simulated PT Patient &mdash; Harrisburg University</p>
  <div class="consent-box">
    <h3>Informed Consent</h3>
    <ol>
      <li><strong>Session recording.</strong> All chat interactions are automatically saved and may be reviewed by course instructors and the development team at Harrisburg University.</li>
      <li><strong>Intended use only.</strong> This AI simulation is designed exclusively for physical therapy (PT) patient roleplay exercises. Do not use it for medical advice, diagnosis, or any purpose outside the assigned educational activity.</li>
      <li><strong>Research and improvement.</strong> Your session data may be used &mdash; in anonymized or identifiable form &mdash; to evaluate system performance and improve the simulation tool as part of ongoing educational research.</li>
      <li><strong>Subject to change.</strong> The AI model, patient characters, interface, and all other system components may be updated or replaced at any time without prior notice.</li>
      <li><strong>Voluntary participation.</strong> Use of this tool is voluntary. You may stop at any time by logging out.</li>
    </ol>
  </div>
  {% if error %}<div class="error">{{ error }}</div>{% endif %}
  <form method="POST">
    <input type="text" name="username" placeholder="Choose a username (min. 3 characters)" autofocus autocomplete="username">
    <input type="password" name="password" placeholder="Choose a password (min. 6 characters)" autocomplete="new-password">
    <input type="password" name="confirm" placeholder="Confirm password" autocomplete="new-password">
    <label class="checkbox-row">
      <input type="checkbox" name="consent" value="yes">
      I have read and understood the informed consent above, and I agree to participate.
    </label>
    <button type="submit">Create Account</button>
  </form>
  <div class="alt-link">Already have an account? <a href="/login">Sign in</a></div>
</div>
</body>
</html>"""


# ── Sessions list page ────────────────────────────────────────────────────────
SESSIONS_HTML = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>My Sessions — Simulated PT Patient</title>
<style>
  * { box-sizing: border-box; margin: 0; padding: 0; }
  body { font-family: system-ui, sans-serif; background: #f0f2f5; }
  header { background: #0C6157; color: white; padding: 12px 20px;
           display: flex; align-items: center; justify-content: space-between; }
  .logo { height: 40px; }
  header h1 { font-size: 1rem; font-weight: 600; }
  .hdr-btn { background: #CBB778; border: none; color: white; border-radius: 8px;
             padding: 6px 14px; cursor: pointer; font-size: 0.85rem;
             text-decoration: none; display: inline-block; }
  .hdr-btn:hover { background: #b5a265; }
  .container { max-width: 900px; margin: 32px auto; padding: 0 20px; }
  h2 { font-size: 1rem; color: #0C6157; margin-bottom: 16px; }
  table { width: 100%; border-collapse: collapse; background: white;
          border-radius: 10px; overflow: hidden; box-shadow: 0 1px 4px rgba(0,0,0,.1); }
  th { background: #0C6157; color: white; padding: 10px 16px; text-align: left; font-size: 0.83rem; }
  td { padding: 10px 16px; font-size: 0.9rem; border-bottom: 1px solid #eee; }
  tr:last-child td { border-bottom: none; }
  tr:hover td { background: #fafafa; }
  .action-link { color: #0C6157; text-decoration: none; font-weight: 500; margin-right: 12px; }
  .action-link:hover { text-decoration: underline; }
  .dl-link { color: #888; text-decoration: none; font-size: 0.82rem; }
  .dl-link:hover { color: #0C6157; text-decoration: underline; }
  .note-dot { color: #CBB778; font-size: 0.9rem; }
  .empty { text-align: center; color: #888; padding: 48px 20px; font-size: 0.95rem;
           background: white; border-radius: 10px; box-shadow: 0 1px 4px rgba(0,0,0,.1); }
</style>
</head>
<body>
<header>
  <img class="logo" src="https://www.arcgis.com/sharing/rest/content/items/088d68905927400bb34449dc1b387446/resources/images/widget_2/1709839675447.png" alt="logo">
  <h1>My Sessions</h1>
  <a href="/" class="hdr-btn">&#8592; Back to Chat</a>
</header>
<div class="container">
  <h2>Sessions for {{ username }}</h2>
  {% if sessions %}
  <table>
    <thead>
      <tr>
        <th>Date &amp; Time</th>
        <th>Character</th>
        <th>Exchanges</th>
        <th>Notes</th>
        <th>Actions</th>
      </tr>
    </thead>
    <tbody>
      {% for s in sessions %}
      <tr>
        <td>{{ s.started_at }}</td>
        <td>{{ s.character }}</td>
        <td>{{ s.exchange_count }}</td>
        <td>{% if s.has_notes %}<span class="note-dot">&#9679;</span>{% endif %}</td>
        <td>
          <a class="action-link" href="/sessions/{{ s.session_id }}">View</a>
          <a class="dl-link" href="/sessions/{{ s.session_id }}/download">&#11015; Excel</a>
        </td>
      </tr>
      {% endfor %}
    </tbody>
  </table>
  {% else %}
  <div class="empty">No sessions yet. Start a chat to create your first session.</div>
  {% endif %}
</div>
</body>
</html>"""


# ── Session detail page ───────────────────────────────────────────────────────
SESSION_DETAIL_HTML = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Session — Simulated PT Patient</title>
<style>
  * { box-sizing: border-box; margin: 0; padding: 0; }
  body { font-family: system-ui, sans-serif; background: #f0f2f5; }
  header { background: #0C6157; color: white; padding: 12px 20px;
           display: flex; align-items: center; justify-content: space-between; gap: 12px; }
  .logo { height: 40px; flex-shrink: 0; }
  header h1 { font-size: 0.95rem; font-weight: 600; }
  .hdr-btn { background: #CBB778; border: none; color: white; border-radius: 8px;
             padding: 6px 14px; cursor: pointer; font-size: 0.85rem;
             text-decoration: none; display: inline-block; white-space: nowrap; }
  .hdr-btn:hover { background: #b5a265; }
  .container { max-width: 760px; margin: 32px auto; padding: 0 20px 40px; }
  .meta { color: #777; font-size: 0.82rem; margin-bottom: 20px; }
  .notes-card { background: #fffdf5; border: 1px solid #e8dcb0; border-left: 4px solid #CBB778;
                border-radius: 8px; padding: 14px 16px; margin-bottom: 26px; }
  .notes-card h3 { font-size: 0.72rem; color: #a08a3c; text-transform: uppercase;
                   letter-spacing: .06em; margin-bottom: 8px; }
  .notes-card p { font-size: 0.88rem; color: #444; white-space: pre-wrap; line-height: 1.55; }
  .exchange { margin-bottom: 20px; }
  .q-wrap { text-align: right; margin-bottom: 6px; }
  .a-wrap { text-align: left; }
  .label { font-size: 0.72rem; color: #888; margin-bottom: 3px; }
  .q { background: #1a73e8; color: white; padding: 10px 14px;
       border-radius: 18px 18px 4px 18px; font-size: 0.9rem;
       display: inline-block; max-width: 80%; text-align: left; }
  .a { background: white; color: #111; padding: 10px 14px;
       border-radius: 18px 18px 18px 4px; font-size: 0.9rem;
       display: inline-block; max-width: 80%; text-align: left;
       box-shadow: 0 1px 2px rgba(0,0,0,.1); white-space: pre-wrap; }
</style>
</head>
<body>
<header>
  <img class="logo" src="https://www.arcgis.com/sharing/rest/content/items/088d68905927400bb34449dc1b387446/resources/images/widget_2/1709839675447.png" alt="logo">
  <h1>{{ data.character }} &mdash; {{ data.started_at }}</h1>
  <div style="display:flex;gap:8px">
    <a href="/sessions/{{ data.session_id }}/download" class="hdr-btn">&#11015; Excel</a>
    <a href="/sessions" class="hdr-btn">&#8592; My Sessions</a>
  </div>
</header>
<div class="container">
  <div class="meta">{{ data.exchange_count }} exchanges &nbsp;&middot;&nbsp; {{ data.username }}</div>
  {% if data.notes %}
  <div class="notes-card">
    <h3>Scrap Notes</h3>
    <p>{{ data.notes }}</p>
  </div>
  {% endif %}
  {% for ex in data.exchanges %}
  <div class="exchange">
    <div class="q-wrap">
      <div class="label">You</div>
      <div class="q">{{ ex.question }}</div>
    </div>
    <div class="a-wrap">
      <div class="label">{{ data.character }}</div>
      <div class="a">{{ ex.response }}</div>
    </div>
  </div>
  {% endfor %}
</div>
</body>
</html>"""


# ── Flask app ─────────────────────────────────────────────────────────────────
app = Flask(__name__)
app.secret_key = secrets.token_hex(32)


def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not session.get("authenticated"):
            return redirect("/login")
        return f(*args, **kwargs)
    return decorated


# ── Auth routes ───────────────────────────────────────────────────────────────
@app.route("/register", methods=["GET", "POST"])
def register():
    error = ""
    if request.method == "POST":
        username = request.form.get("username", "").strip().lower()
        password = request.form.get("password", "")
        confirm  = request.form.get("confirm", "")
        consent  = request.form.get("consent", "")
        if not username or not password:
            error = "All fields are required."
        elif " " in username or len(username) < 3:
            error = "Username must be at least 3 characters with no spaces."
        elif len(password) < 6:
            error = "Password must be at least 6 characters."
        elif password != confirm:
            error = "Passwords do not match."
        elif not consent:
            error = "You must accept the informed consent to register."
        else:
            users = load_users()
            if username in users:
                error = "Username already taken. Choose another."
            else:
                # First account created becomes the developer account.
                role = "developer" if not users else "student"
                users[username] = {
                    "password":   generate_password_hash(password),
                    "role":       role,
                    "created_at": datetime.now(timezone.utc).isoformat(),
                }
                save_users(users)
                session["authenticated"] = True
                session["username"] = username
                return redirect("/")
    return render_template_string(REGISTER_HTML, error=error)


@app.route("/login", methods=["GET", "POST"])
def login():
    error = ""
    if request.method == "POST":
        username = request.form.get("username", "").strip().lower()
        password = request.form.get("password", "")
        record = load_users().get(username)
        if record and check_password_hash(record["password"], password):
            session["authenticated"] = True
            session["username"] = username
            return redirect("/")
        error = "Incorrect username or password."
    return render_template_string(LOGIN_HTML, error=error)


@app.route("/logout")
def logout():
    session.clear()
    return redirect("/login")


# ── App routes ────────────────────────────────────────────────────────────────
@app.route("/")
@login_required
def index():
    return render_template_string(HTML,
                                  username=session.get("username", ""),
                                  role=current_role())


@app.route("/characters")
@login_required
def characters():
    meta = load_meta()
    return jsonify({"characters": [
        {"name": n, "gender": meta.get(n, {}).get("gender", "female")}
        for n in character_names()
    ]})


@app.route("/character/<name>")
@login_required
def character(name):
    text = read_character(name)
    if not text:
        return jsonify({"error": "not found"}), 404
    return jsonify({"context": text})


@app.route("/avatar/<name>")
@login_required
def avatar(name):
    if name in character_names():
        for ext in ("png", "jpg", "jpeg", "webp"):
            candidate = AVATARS_DIR / f"{name}.{ext}"
            if candidate.exists():
                return send_file(candidate)
    return Response(build_avatar_svg(name), mimetype="image/svg+xml")


@app.route("/chat", methods=["POST"])
@login_required
def chat():
    data           = request.get_json()
    user_message   = data["message"]
    history        = data.get("history", [])
    session_id     = data.get("session_id", "")
    character_name = data.get("character", "")
    username       = session.get("username", "anonymous")

    # The system context is resolved server-side from the character file, so a
    # student cannot alter the prompt from the browser. Developers may override.
    override = data.get("override_context", "").strip()
    system_context = override if (override and current_role() == "developer") \
        else read_character(character_name)

    messages = []
    if system_context:
        messages.append({"role": "system", "content": system_context})
    for past_user, past_bot in history:
        messages.append({"role": "user",      "content": past_user})
        messages.append({"role": "assistant", "content": past_bot})
    messages.append({"role": "user", "content": user_message})

    try:
        from openai import OpenAI
        client = OpenAI(base_url=API_URL, api_key=API_KEY)
        resp = client.chat.completions.create(model=MODEL_NAME, messages=messages)
        reply = resp.choices[0].message.content
    except Exception as e:
        reply = f"Connection error: {e}"

    full_history = history + [[user_message, reply]]
    if session_id:
        save_session(session_id, username, character_name, full_history)

    return jsonify({"reply": reply})


@app.route("/session-notes", methods=["POST"])
@login_required
def session_notes():
    data       = request.get_json()
    session_id = data.get("session_id", "")
    if not session_id:
        return jsonify({"ok": False}), 400

    username = session.get("username", "")
    user_dir = SESSIONS_DIR / username
    user_dir.mkdir(exist_ok=True)
    path = _session_path(username, session_id)
    now  = datetime.now(timezone.utc).isoformat()

    record = {}
    if path.exists():
        try:
            record = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            record = {}

    record.setdefault("session_id", session_id)
    record.setdefault("username", username)
    record.setdefault("character", data.get("character", ""))
    record.setdefault("started_at", now)
    record.setdefault("exchanges", [])
    record.setdefault("exchange_count", len(record.get("exchanges", [])))
    record["notes"]        = data.get("notes", "")
    record["last_updated"] = now

    path.write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")
    return jsonify({"ok": True})


@app.route("/tts", methods=["POST"])
@login_required
def tts():
    """Synthesize a patient reply with edge-tts (free Microsoft neural voices).

    Only the patient's words are sent — never anything the student typed.
    Returns 503 when unavailable so the browser falls back to its own voices.
    """
    try:
        import asyncio
        import edge_tts
    except ImportError:
        return "edge-tts not installed", 503

    data = request.get_json() or {}
    text = (data.get("text") or "").strip()
    if not text:
        return "", 400
    text = text[:MAX_TTS_CHARS]

    meta  = load_meta().get(data.get("character", ""), {})
    voice = meta.get("voice") or DEFAULT_VOICES.get(meta.get("gender", "female"),
                                                    DEFAULT_VOICES["female"])

    async def synth():
        buf = io.BytesIO()
        async for chunk in edge_tts.Communicate(text, voice).stream():
            if chunk["type"] == "audio":
                buf.write(chunk["data"])
        return buf

    try:
        buf = asyncio.run(synth())
    except Exception as e:
        app.logger.warning("edge-tts failed: %s", e)
        return "tts failed", 503

    if not buf.getbuffer().nbytes:
        return "empty audio", 503

    buf.seek(0)
    return send_file(buf, mimetype="audio/mpeg")


@app.route("/test-questions")
@login_required
def test_questions():
    path = CHARS_DIR / "test_questions.txt"
    if not path.exists():
        return jsonify({"questions": []})
    questions = [l.strip() for l in path.read_text(encoding="utf-8").splitlines() if l.strip()]
    return jsonify({"questions": questions})


@app.route("/sessions")
@login_required
def sessions_list():
    username = session.get("username", "")
    return render_template_string(SESSIONS_HTML,
                                  username=username,
                                  sessions=get_user_sessions(username))


@app.route("/sessions/<sid>")
@login_required
def session_detail(sid):
    data = get_session_data(session.get("username", ""), sid)
    if data is None:
        return "Session not found.", 404
    return render_template_string(SESSION_DETAIL_HTML, data=data)


@app.route("/sessions/<sid>/download")
@login_required
def session_download(sid):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        return "openpyxl not installed. Run: venv/bin/pip install openpyxl", 500

    username = session.get("username", "")
    data = get_session_data(username, sid)
    if data is None:
        return "Session not found.", 404

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Session"

    character  = data.get("character", "Patient")
    started_at = data.get("started_at", "")

    ws.append(["Session ID",  data.get("session_id", "")])
    ws.append(["User",        username])
    ws.append(["Character",   character])
    ws.append(["Date",        started_at])
    ws.append(["Scrap notes", data.get("notes", "")])
    ws.append([])

    header_row = ws.max_row + 1
    ws.append(["#", "Author", "Message"])
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0C6157")

    for i, ex in enumerate(data.get("exchanges", []), 1):
        ws.append([i, "Student", ex.get("question", "")])
        ws.append([i, character, ex.get("response", "")])

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 90
    ws.cell(row=5, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=header_row + 1):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True,
                     download_name=f"session_{character}_{started_at[:10]}.xlsx")


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=80, debug=False)
